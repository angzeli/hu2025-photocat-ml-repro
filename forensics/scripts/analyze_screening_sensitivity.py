#!/usr/bin/env python3
"""Bounded Fig. 1 threshold and SI Table 3 audit; no fitting or source mirrors.

Read only the released 34 screening points, 43 measured comparisons and two
SI Table 3 pages. Full coordinates and measurements stay in memory. Tracked
outputs are 49 threshold summaries, aggregate validation counts and a small
published identity set; no grid cell discloses a selected source-row list.
"""

import csv
import json
from pathlib import Path
import re

import numpy as np
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
FIG1 = "orginal/41929_2025_1291_MOESM11_ESM.xlsx"
SI = "orginal/41929_2025_1291_MOESM1_ESM.pdf"
J_THRESHOLDS = (20., 30., 40., 50., 60., 70., 80.)
JSTAR_THRESHOLDS = (0.002, 0.003, 0.005, 0.01, 0.015, 0.02, 0.03)


def read_fig1():
    """Return source arrays in memory, retaining observed row order and IDs."""
    workbook = load_workbook(ROOT / FIG1, read_only=True, data_only=True)
    try:
        points = []
        for source_row, values in enumerate(workbook["Fig.1d"].values, 1):
            if source_row == 1 or all(value is None for value in values):
                continue
            assert len(values) == 2 and all(isinstance(value, (float, int)) for value in values)
            points.append((source_row, *values))
        measured = {}
        for cat, ps, ton, selectivity in workbook["Fig.1e"].iter_rows(min_row=2, values_only=True):
            if cat is None:
                continue
            identity = (int(cat.split()[-1]), int(ps.split()[-1]))
            assert identity not in measured
            measured[identity] = (float(ton), float(selectivity))
        return np.array(points, dtype=float), measured
    finally:
        workbook.close()


def read_table3():
    """Parse the explicit Filtered/Ruled-out boundary, not a performance ordering."""
    reader = PdfReader(ROOT / SI)
    text = "\n".join(reader.pages[i].extract_text() for i in (76, 77))
    filtered_text, ruled_out_text = re.split(r"Ruled-\s*out", text.split("Filtered", 1)[1], maxsplit=1)
    value = r"(?:\d+(?:\.\d+)?|−)"
    pattern = re.compile(rf"CAT\s+(\d+)\s+PS\s+(\d+)\s+({value})\s+({value})\s+({value})\s+({value})")
    groups = {}
    for name, block in (("filtered", filtered_text), ("ruled_out", ruled_out_text)):
        rows = {}
        for cat, ps, ton, tof, selectivity, time in pattern.findall(block):
            identity = (int(cat), int(ps))
            assert identity not in rows
            numbers = (ton, tof, selectivity, time)
            assert all(v == "−" for v in numbers) or all(v != "−" for v in numbers)
            rows[identity] = None if ton == "−" else (float(ton), float(selectivity))
        groups[name] = rows
    assert not (set(groups["filtered"]) & set(groups["ruled_out"]))
    return groups


def same_set_interval(values, other_pass, selected):
    """One-axis threshold interval [lower, upper) preserving a strict-> set."""
    assert np.any(selected) and np.any(other_pass & ~selected)
    return [float(values[other_pass & ~selected].max()), float(values[selected].min())]


def main():
    points, measured = read_fig1()
    assert points.shape == (34, 3) and len(measured) == 43
    assert np.isfinite(points).all() and np.all(points[:, 1:] > 0)
    j, jstar = points[:, 1], points[:, 2]
    baseline = (j > 50) & (jstar > 0.01)
    assert int(baseline.sum()) == 7
    intervals = {
        "J_with_Jstar_fixed_0p01": same_set_interval(j, jstar > 0.01, baseline),
        "Jstar_with_J_fixed_50": same_set_interval(jstar, j > 50, baseline),
        "interval_convention": "lower inclusive, upper exclusive; each conditional on the other stated fixed threshold",
    }
    old = json.loads((ROOT / "forensics/outputs/screening_analysis.json").read_text())
    for key in ("J_with_Jstar_fixed_0p01", "Jstar_with_J_fixed_50"):
        assert intervals[key] == old["fig1d"]["same_selection_threshold_intervals"][key]

    grid = []
    for j_cutoff in J_THRESHOLDS:
        for js_cutoff in JSTAR_THRESHOLDS:
            selected = (j > j_cutoff) & (jstar > js_cutoff)
            overlap = int(np.sum(selected & baseline))
            union = int(np.sum(selected | baseline))
            grid.append({
                "J_threshold_exclusive": j_cutoff,
                "Jstar_threshold_exclusive": js_cutoff,
                "supplied_point_count": len(j),
                "selected_count": int(selected.sum()),
                "overlap_with_probe_seven": overlap,
                "added_to_probe_count": int(np.sum(selected & ~baseline)),
                "removed_from_probe_count": int(np.sum(baseline & ~selected)),
                "jaccard_with_probe_seven": overlap / union,
                "same_seven_point_set": bool(np.array_equal(selected, baseline)),
            })
    assert {6, 7, 8}.issubset({row["selected_count"] for row in grid})

    groups = read_table3()
    filtered = groups["filtered"]
    ruled_out = groups["ruled_out"]
    assert len(filtered) == 7 and len(ruled_out) == 37
    unmeasured = {key for key, values in filtered.items() if values is None}
    assert unmeasured == {(50, 41)}
    filtered_measured = {key: values for key, values in filtered.items() if values is not None}
    expected_measured = filtered_measured | ruled_out
    assert len(filtered_measured) == 6 and expected_measured == measured
    # This fixed diagnostic is inherited from Phase 0, not fitted to these outcomes.
    conditional_good = {key for key, (ton, selectivity) in measured.items()
                        if ton > 1370 and selectivity > 72}
    predicted = set(filtered_measured)
    tp, fp = len(predicted & conditional_good), len(predicted - conditional_good)
    fn, tn = len(conditional_good - predicted), len(set(measured) - predicted - conditional_good)
    assert (tp, fp, fn, tn) == (6, 0, 1, 36)
    assert conditional_good - predicted == {(61, 33)}
    ton = np.array([values[0] for values in measured.values()])
    selectivity = np.array([values[1] for values in measured.values()])
    good_mask = (ton > 1370) & (selectivity > 72)
    good_intervals = {
        "TON_with_selectivity_fixed_72": same_set_interval(ton, selectivity > 72, good_mask),
        "selectivity_with_TON_fixed_1370": same_set_interval(selectivity, ton > 1370, good_mask),
        "interval_convention": "lower inclusive, upper exclusive; conditional one-axis intervals, not a joint rectangle",
    }
    for key in ("TON_with_selectivity_fixed_72", "selectivity_with_TON_fixed_1370"):
        assert good_intervals[key] == old["fig1e"]["conditional_good_threshold_intervals"][key]

    # Phase-0 figure/text inference establishes this identity; numeric recheck
    # confirms its unique highest-J* location, not a spreadsheet identity label.
    star = points[points[:, 0] == 2]
    assert len(star) == 1 and np.sum(jstar == star[0, 2]) == 1 and star[0, 2] == jstar.max()
    assert baseline[points[:, 0] == 2].item()
    regions = []
    for j_cutoff in J_THRESHOLDS:
        rows = [row for row in grid if row["J_threshold_exclusive"] == j_cutoff]
        regions.append({
            "J_threshold_exclusive": j_cutoff,
            "Jstar_grid_values_with_count_seven": [row["Jstar_threshold_exclusive"] for row in rows if row["selected_count"] == 7],
            "Jstar_grid_values_with_same_seven_point_set": [row["Jstar_threshold_exclusive"] for row in rows if row["same_seven_point_set"]],
        })
    result = {
        "classification": "COMPUTED FROM RELEASED DATA",
        "sources": {"workbook": FIG1, "sheets": ["Fig.1d", "Fig.1e"], "SI": SI, "SI_PDF_pages": [77, 78]},
        "boundary": "49 compact grid rows; no source coordinates, measurement columns or per-cell source-row sets emitted; no fitting",
        "screening": {
            "supplied_points": len(j), "all_positive": True, "full_retained_universe": 180,
            "missing_scores_not_reconstructed": 146,
            "probe": {"J_gt": 50, "Jstar_gt": 0.01, "count": int(baseline.sum()), "status": "compatibility probe, not uniquely identified author thresholds"},
            "same_set_conditional_intervals": intervals,
            "phase0_intervals_reproduced": True,
            "grid": {
                "J_thresholds": J_THRESHOLDS, "Jstar_thresholds": JSTAR_THRESHOLDS, "cells": len(grid),
                "count_distribution": {str(n): sum(row["selected_count"] == n for row in grid) for n in sorted({row["selected_count"] for row in grid})},
                "cells_with_seven_points": sum(row["selected_count"] == 7 for row in grid),
                "cells_with_same_seven_point_set": sum(row["same_seven_point_set"] for row in grid),
                "seven_count_different_set_cells": sum(row["selected_count"] == 7 and not row["same_seven_point_set"] for row in grid),
                "seven_point_regions_on_sampled_grid": regions,
                "scope": "sampled threshold coordinates only; not an exhaustive continuous-region map or experimental outcome optimization",
            },
            "selected_identity_set": [f"CAT{cat}_PS{ps}" for cat, ps in sorted(filtered)],
            "identity_set_evidence": "DIRECTLY OBSERVED: SI Table 3 filtered group",
            "star_mapping": {"source_row": 2, "identity": "CAT1_PS1", "classification": "STRONG INFERENCE", "evidence": "Phase-0 article Fig1 star/best-system cross-source anchor; unique largest-Jstar row rechecked"},
            "remaining_selected_point_identities": "six unresolved; no ordering imposed from Table 3",
        },
        "experimental_validation": {
            "SI_total_including_unmeasured": len(filtered) + len(ruled_out),
            "measured_filtered": len(filtered_measured), "measured_ruled_out": len(ruled_out),
            "workbook_SI_measured_value_discrepancies": 0,
            "unmeasured_selected_identity": "CAT50_PS41",
            "conditional_ruled_out_success": "CAT61_PS33",
            "declared_arithmetic": {"precision_numerator": 6, "precision_denominator": 6, "precision": 6 / 6,
                                    "recall_numerator": 6, "recall_denominator": 7, "recall": 6 / 7},
            "unmeasured_denominator_reason": "CAT50/PS41 was not synthesized/measured; it is not an observed false positive and is outside measured precision",
            "conditional_good_probe": {"rule": "TON_CO > 1370 AND CO selectivity (%) > 72", "tp": tp, "fp": fp, "fn": fn, "tn": tn,
                                       "precision": tp / (tp + fp), "recall": tp / (tp + fn), "status": "compatible but not uniquely recovered author classifier"},
            "conditional_good_same_set_intervals": good_intervals,
            "scope": "43 experimentally measured identities only; no recall claim over 180 or 3444 systems",
        },
    }
    output_csv = ROOT / "data/derived/screening_sensitivity.csv"
    with output_csv.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(grid[0]))
        writer.writeheader()
        writer.writerows(grid)
    output_json = ROOT / "forensics/outputs/phase2_screening.json"
    output_json.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"screening_points": len(j), "probe_count": int(baseline.sum()),
                      "grid_count_distribution": result["screening"]["grid"]["count_distribution"],
                      "same_seven_grid_cells": result["screening"]["grid"]["cells_with_same_seven_point_set"],
                      "count_seven_different_set_cells": result["screening"]["grid"]["seven_count_different_set_cells"],
                      "conditional_confusion": [tp, fp, fn, tn], "source_values_emitted": False}, indent=2))


if __name__ == "__main__":
    main()
